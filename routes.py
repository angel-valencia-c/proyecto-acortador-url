from flask import Blueprint, request, redirect, jsonify, render_template, url_for, Response, send_file
import datetime
import string
import random
import sqlite3
import io
from database import get_db_connection

# Al inicio de routes.py — agrega estos imports
from auth import login_required, handle_login, handle_logout
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    openpyxl = None 

bp = Blueprint('main', __name__)

@bp.route('/login', methods=['GET', 'POST'])
def login():
    return handle_login()

@bp.route('/logout')
def logout():
    return handle_logout()


def generate_short_id(length=6):
    chars = string.ascii_letters + string.digits
    return ''.join(random.choice(chars) for _ in range(length))


@bp.route('/')
def home():
    return render_template('index.html')

@bp.route('/crear', methods=['POST'])
def crear_url():
    data = request.get_json()
    original_url = data.get('original_url')
    
    if not original_url:
        return jsonify({'error': 'URL requerida'}), 400
        
    short_id = generate_short_id()
    
    conn = get_db_connection()
    conn.execute("INSERT INTO urls (short_id, original_url) VALUES (?, ?)", (short_id, original_url))
    conn.commit()
    conn.close()
    
    return jsonify({'short_id': short_id})


@bp.route('/<short_id>')
def redireccionar(short_id):
    if short_id in ('favicon.ico', 'robots.txt'):
        return "", 404

    conn = get_db_connection()
    url_row = conn.execute(
        "SELECT original_url, is_active FROM urls WHERE short_id = ?", 
        (short_id,)
    ).fetchone()

    if not url_row:
        conn.close()
        return "Enlace no encontrado", 404

    if not url_row['is_active']:
        conn.close()
        return "Este enlace ha sido desactivado", 410

    original_url = url_row['original_url']

    # Capturar todos los parámetros
    ua_string = request.headers.get('User-Agent', '')
    
    # Device type simple sin librería externa
    ua_lower = ua_string.lower()
    if any(x in ua_lower for x in ['mobile', 'android', 'iphone']):
        device_type = 'mobile'
    elif 'tablet' in ua_lower or 'ipad' in ua_lower:
        device_type = 'tablet'
    else:
        device_type = 'desktop'

    # Browser simple
    if 'edg' in ua_lower:
        browser = 'Edge'
    elif 'chrome' in ua_lower:
        browser = 'Chrome'
    elif 'firefox' in ua_lower:
        browser = 'Firefox'
    elif 'safari' in ua_lower:
        browser = 'Safari'
    else:
        browser = 'Otro'

    conn.execute('''
        INSERT INTO visits (
            short_id, ip_address, user_agent, referer,
            utm_source, utm_medium, utm_campaign, utm_term, utm_content,
            gclid, fbclid, ttclid, msclkid,
            device_type, browser, country, timestamp
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        short_id,
        request.remote_addr,
        ua_string,
        request.referrer,
        request.args.get('utm_source'),
        request.args.get('utm_medium'),
        request.args.get('utm_campaign'),
        request.args.get('utm_term'),
        request.args.get('utm_content'),
        request.args.get('gclid'),
        request.args.get('fbclid'),
        request.args.get('ttclid'),
        request.args.get('msclkid'),
        device_type,
        browser,
        None,  # country — lo conectamos después con ip-api
        datetime.datetime.now()
    ))

    conn.commit()
    conn.close()
    return redirect(original_url, code=302)
    if short_id == 'favicon.ico':
        return "", 404

    conn = get_db_connection()
    url_row = conn.execute("SELECT original_url FROM urls WHERE short_id = ?", (short_id,)).fetchone()
    
    if url_row:
        original_url = url_row['original_url']
        
        utm_source = request.args.get('utm_source')
        utm_medium = request.args.get('utm_medium')
        utm_campaign = request.args.get('utm_campaign')
        
        conn.execute('''
            INSERT INTO visits (short_id, ip_address, user_agent, referer, utm_source, utm_medium, utm_campaign, timestamp)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        ''', (short_id, request.remote_addr, request.headers.get('User-Agent'), request.headers.get('Referer'), 
              utm_source, utm_medium, utm_campaign, datetime.datetime.now()))
        conn.commit()
        conn.close()
        
        if request.query_string:
            params = request.query_string.decode("utf-8")
            final_url = f"{original_url}?{params}" if "?" not in original_url else f"{original_url}&{params}"
            return redirect(final_url)
            
        return redirect(original_url)
        
    conn.close()
    return "<h1>404 - Enlace no encontrado</h1>", 404


@bp.route('/reportes')
@login_required
def dashboard():
    conn = get_db_connection()

    # Visitas con join a urls
    visitas = conn.execute('''
        SELECT v.*, u.original_url, u.is_active
        FROM visits v
        LEFT JOIN urls u ON v.short_id = u.short_id
        ORDER BY v.timestamp DESC
    ''').fetchall()

    # Métricas para las cards
    total_clics = conn.execute('SELECT COUNT(*) FROM visits').fetchone()[0]
    total_links = conn.execute('SELECT COUNT(*) FROM urls WHERE is_active = 1').fetchone()[0]

    # Clics por día (últimos 7 días)
    clics_por_dia = conn.execute('''
        SELECT DATE(timestamp) as dia, COUNT(*) as total
        FROM visits
        WHERE timestamp >= DATE('now', '-7 days')
        GROUP BY DATE(timestamp)
        ORDER BY dia ASC
    ''').fetchall()

    # Device breakdown
    devices = conn.execute('''
        SELECT device_type, COUNT(*) as total
        FROM visits
        WHERE device_type IS NOT NULL
        GROUP BY device_type
    ''').fetchall()

    # Top UTM sources
    top_sources = conn.execute('''
        SELECT utm_source, COUNT(*) as total
        FROM visits
        WHERE utm_source IS NOT NULL
        GROUP BY utm_source
        ORDER BY total DESC
        LIMIT 5
    ''').fetchall()

    # Top browsers
    browsers = conn.execute('''
        SELECT browser, COUNT(*) as total
        FROM visits
        WHERE browser IS NOT NULL
        GROUP BY browser
        ORDER BY total DESC
    ''').fetchall()

    conn.close()

    return render_template('dashboard.html',
        visitas=visitas,
        total_clics=total_clics,
        total_links=total_links,
        clics_por_dia=[dict(r) for r in clics_por_dia],
        devices=[dict(r) for r in devices],
        top_sources=[dict(r) for r in top_sources],
        browsers=[dict(r) for r in browsers],
    )

@bp.route('/editar', methods=['POST'])
def editar_url():
    data = request.get_json()
    short_id = data.get('short_id')
    new_url = data.get('new_url')
    
    if not short_id or not new_url:
        return jsonify({'error': 'Faltan datos'}), 400
        
    conn = get_db_connection()
    conn.execute("UPDATE urls SET original_url = ? WHERE short_id = ?", (new_url, short_id))
    conn.commit()
    conn.close()
    
    return jsonify({'message': 'Actualizado correctamente'})

@bp.route('/eliminar/<short_id>', methods=['DELETE'])
def eliminar_url(short_id):
    conn = get_db_connection()
    conn.execute("DELETE FROM visits WHERE short_id = ?", (short_id,))
    conn.execute("DELETE FROM urls WHERE short_id = ?", (short_id,))
    conn.commit()
    conn.close()
    return jsonify({'message': 'Eliminado correctamente'})

@bp.route('/toggle/<short_id>', methods=['POST'])
@login_required
def toggle_link(short_id):
    """Activa o desactiva un enlace corto."""
    conn = get_db_connection()
    current = conn.execute(
        "SELECT is_active FROM urls WHERE short_id = ?", (short_id,)
    ).fetchone()
    
    if not current:
        conn.close()
        return jsonify({'error': 'No encontrado'}), 404
    
    new_state = 0 if current['is_active'] else 1
    conn.execute(
        "UPDATE urls SET is_active = ?, updated_at = datetime('now') WHERE short_id = ?",
        (new_state, short_id)
    )
    conn.commit()
    conn.close()
    return jsonify({'is_active': new_state})

@bp.route('/exportar-excel')
def exportar_excel():
    if not openpyxl:
        return "Error: La librería 'openpyxl' no está instalada. Ejecuta 'pip install openpyxl'", 500

    conn = get_db_connection()
    query = """
        SELECT v.timestamp, v.short_id, u.original_url, v.utm_source, v.utm_medium, 
               v.utm_campaign, v.ip_address, v.user_agent, v.country
        FROM visits v 
        JOIN urls u ON v.short_id = u.short_id 
        ORDER BY v.id DESC
    """
    visitas = conn.execute(query).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trafico"


    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['Fecha', 'ID Corto', 'URL Destino', 'Fuente', 'Medio', 'Campaña', 'IP', 'Navegador', 'País']
    
    
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header_title)
        cell.fill = header_fill
        cell.font = header_font


    for row_num, visita in enumerate(visitas, 2):
        ws.cell(row=row_num, column=1, value=visita['timestamp'])
        ws.cell(row=row_num, column=2, value=visita['short_id'])
        ws.cell(row=row_num, column=3, value=visita['original_url'])
        ws.cell(row=row_num, column=4, value=visita['utm_source'] or '-')
        ws.cell(row=row_num, column=5, value=visita['utm_medium'] or '-')
        ws.cell(row=row_num, column=6, value=visita['utm_campaign'] or '-')
        ws.cell(row=row_num, column=7, value=visita['ip_address'])
        ws.cell(row=row_num, column=8, value=visita['user_agent'])
       
        country = visita['country'] if 'country' in visita.keys() else 'Desconocido'
        ws.cell(row=row_num, column=9, value=country or 'Desconocido')


    ws.column_dimensions['A'].width = 22 # Fecha
    ws.column_dimensions['C'].width = 35 # URL
    ws.column_dimensions['H'].width = 25 # User Agent

  
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='reporte_trafico.xlsx'
    )